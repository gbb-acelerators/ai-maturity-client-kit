#!/usr/bin/env python3
"""Write the auditable v2 scoring workbook from responses.json.

Every score in the workbook is an Excel formula over the raw answers,
so a reviewer can trace each number. The engine value sits next to each
formula as a cross-check (column "Engine"); the two must match.

Sheets:
    README      method and how to read the workbook
    Answers     one row per question, one column per respondent
                (0-4, "NA", or blank when not answered)
    Questions   question score = AVERAGE of the respondent values
    Dimensions  dimension score = AVERAGE of its question scores,
                weight, target, gap, priority
    Overall     weighted mean of the dimensions that have a score

Usage:
    python3 scripts/fill_workbook_v2.py [--responses X.json] [--out DIR]
"""
from __future__ import annotations

import argparse
import datetime
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from kit_files import responses_file  # noqa: E402


def level_formula(ref: str, bands: list[dict]) -> str:
    expr = f'"{bands[-1]["level"]}"'
    for band in reversed(bands[:-1]):
        expr = (f'IF(ROUND({ref},9)<{band["max"]},"{band["level"]}",'
                f'{expr})')
    return f'=IF({ref}="","",{expr})'


def priority_formula(ref: str, cuts: dict) -> str:
    v = f"ROUND({ref},9)"
    return (f'=IF({ref}="","",IF({v}>={cuts["P0"]},"P0",'
            f'IF({v}>={cuts["P1"]},"P1",IF({v}>={cuts["P2"]},"P2",'
            f'"P3"))))')


def build(fw: dict, respostas: dict, scores: dict | None, loc: str):
    import engine_v2 as v2
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="E5F6FD")
    people = v2.respondents_of(respostas)
    sc = fw["scoring"]
    bands = fw["level_bands"]
    targets = respostas.get("target_overrides") or {}
    weights = respostas.get("dimension_weights") or {}
    engine_q = {q["id"]: q["score"] for q in (scores or {}).get(
        "questions", [])}
    engine_d = {d["id"]: d["score"] for d in (scores or {}).get(
        "dimensions", [])}

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    lines = [
        "AI Maturity Assessment v2: auditable scoring workbook",
        f"Framework {fw['version']} · generated "
        f"{datetime.date.today().isoformat()}",
        "",
        "Answers: 0 to 4 = L0 to L4; NA = not applicable (excluded); "
        "blank = not answered (excluded).",
        "Question score = AVERAGE of the respondent values.",
        "Dimension score = AVERAGE of its question scores; blank when "
        "every question is blank.",
        "Overall = SUMPRODUCT(weight, score) / sum of the weights of the "
        "dimensions that have a score.",
        f"Gap = target (default {sc['default_target']}) minus score, "
        "never below 0. Priority score = weight × gap.",
        "Bands: " + ", ".join(
            f"{b['level']} [{b['min']}, {b['max']}"
            + ("]" if i == len(bands) - 1 else ")")
            for i, b in enumerate(bands)),
        "The Engine column holds the value from output/scores.json; the "
        "formula next to it must match.",
    ]
    for i, text in enumerate(lines, start=1):
        readme.cell(row=i, column=1, value=text)
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 110

    ans = wb.create_sheet("Answers")
    ans.append(["Question", "Dimension", "Title"]
               + [p["id"] for p in people])
    first_col = 4
    last_col = first_col + len(people) - 1
    qrow: dict[str, int] = {}
    for d in fw["dimensions"]:
        for q in d["questions"]:
            row = [q["id"], d["id"], v2.tr(q["title"], loc)]
            for p in people:
                a = p["answers"].get(q["id"])
                if a is None:
                    row.append(None)
                elif a.get("level") is None:
                    row.append("NA")
                else:
                    row.append(int(a["level"]))
            ans.append(row)
            qrow[q["id"]] = ans.max_row
    a_first = get_column_letter(first_col)
    a_last = get_column_letter(max(last_col, first_col))

    qs = wb.create_sheet("Questions")
    qs.append(["Question", "Dimension", "Title", "Score", "Level", "n",
               "NA", "Engine"])
    qs_row: dict[str, int] = {}
    for d in fw["dimensions"]:
        for q in d["questions"]:
            r = qrow[q["id"]]
            rng = f"Answers!{a_first}{r}:{a_last}{r}"
            n = qs.max_row + 1
            qs.append([q["id"], d["id"], v2.tr(q["title"], loc),
                       f'=IF(COUNT({rng})=0,"",AVERAGE({rng}))',
                       level_formula(f"D{n}", bands),
                       f"=COUNT({rng})", f'=COUNTIF({rng},"NA")',
                       engine_q.get(q["id"])])
            qs_row[q["id"]] = n

    ds = wb.create_sheet("Dimensions")
    ds.append(["Dimension", "Name", "Weight", "Score", "Level", "Target",
               "Gap", "Priority score", "Priority", "Engine",
               "Weight × score", "Counted weight"])
    for d in fw["dimensions"]:
        rows = [qs_row[q["id"]] for q in d["questions"]]
        rng = f"Questions!D{min(rows)}:D{max(rows)}"
        n = ds.max_row + 1
        ds.append([
            d["id"], v2.tr(d["name"], loc),
            float(weights.get(d["id"], 1.0)),
            f'=IF(COUNT({rng})=0,"",AVERAGE({rng}))',
            level_formula(f"D{n}", bands),
            float(targets.get(d["id"], sc["default_target"])),
            f'=IF(D{n}="","",MAX(0,F{n}-D{n}))',
            f'=IF(G{n}="","",C{n}*G{n})',
            priority_formula(f"H{n}", sc["priority_cuts"]),
            engine_d.get(d["id"]),
            f'=IF(D{n}="","",C{n}*D{n})',
            f'=IF(D{n}="","",C{n})',
        ])
    last = ds.max_row

    ov = wb.create_sheet("Overall")
    ov.append(["Metric", "Value", "Engine"])
    ov.append(["Overall score",
               f'=IFERROR(SUM(Dimensions!K2:K{last})/'
               f'SUM(Dimensions!L2:L{last}),"")',
               (scores or {}).get("overall", {}).get("score")])
    ov.append(["Overall level", level_formula("B2", bands),
               (scores or {}).get("overall", {}).get("level")])
    ov.append(["Questions with a score",
               f"=COUNT(Questions!D2:D{qs.max_row})",
               (scores or {}).get("threshold", {}).get("answered")])
    ov.append(["Coverage status",
               f'=IF(B4>={sc["coverage_ok_min"]},"OK",IF(B4>='
               f'{sc["coverage_warning_min"]},"WARNING","BLOCKED"))',
               (scores or {}).get("threshold", {}).get("status")])
    ov.append(["Respondents", len(people),
               (scores or {}).get("metadata", {}).get("respondents")])

    for ws in (ans, qs, ds, ov):
        for cell in ws[1]:
            cell.font = bold
            cell.fill = head
        ws.freeze_panes = "B2"
        ws.column_dimensions["A"].width = 14
        if ws.max_column >= 3:
            ws.column_dimensions["C"].width = 44
    ov.column_dimensions["B"].width = 16
    for ws in (qs, ds, ov):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"
        for col in ("B", "D"):
            for cell in ws[col][1:]:
                if isinstance(cell.value, str) and cell.value.startswith(
                        "="):
                    cell.number_format = "0.00"
    wb.calculation.fullCalcOnLoad = True
    return wb


def run(args) -> int:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("✗ openpyxl is required: make install-deps", file=sys.stderr)
        return 1
    src = Path(args.respostas)
    respostas = json.loads(src.read_text(encoding="utf-8"))
    fw = json.loads((ROOT / "framework.v2.json").read_text("utf-8"))
    out_dir = Path(args.out)
    scores_path = out_dir / "scores.json"
    scores = None
    if scores_path.exists():
        data = json.loads(scores_path.read_text(encoding="utf-8"))
        if "dimensions" in data:
            scores = data
    import assessment_engine as engine
    import engine_v2 as v2
    try:
        v2.read_answers(v2.respondents_of(respostas), v2.question_index(fw))
    except v2.InputErrorV2 as exc:
        print(f"✗ {exc}", file=sys.stderr)
        return 1
    wb = build(fw, respostas, scores, engine.locale_of(respostas))
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"scoring-v2-{datetime.date.today().isoformat()}.xlsx"
    wb.save(out)
    print(f"✓ Workbook (v2): {out}")
    if scores is None:
        print("• Engine column is empty: run `make scores` first to "
              "cross-check.")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--responses", "--respostas", dest="respostas",
                    default=None,
                    help="input file (default: responses.json)")
    ap.add_argument("--out", default=str(ROOT / "output"))
    args = ap.parse_args()
    if not args.respostas:
        args.respostas = str(responses_file(ROOT))
    return run(args)


if __name__ == "__main__":
    sys.exit(main())

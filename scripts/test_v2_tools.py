"""Tests for the v2 workbook, round comparison and report payload."""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "relatorios" / "scripts"))

import assessment_engine as engine  # noqa: E402
import compare_rounds  # noqa: E402

MOCK = ROOT / "respostas.v2.json.example"


class V2ToolsTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)
        self.kit = self.tmp / "kit"
        self.out = self.kit / "saida"
        self.out.mkdir(parents=True)
        shutil.copy(ROOT / "framework.v2.json", self.kit)
        shutil.copy(MOCK, self.kit / "respostas.json")
        self.assertEqual(
            engine.run("all", self.kit / "respostas.json", self.out), 0)
        self.scores = json.loads((self.out / "scores.json").read_text())

    def test_workbook_carries_engine_values(self) -> None:
        import openpyxl

        import fill_workbook_v2
        args = SimpleNamespace(respostas=str(self.kit / "respostas.json"),
                               out=str(self.out))
        self.assertEqual(fill_workbook_v2.run(args), 0)
        path = next(self.out.glob("pontuacao-v2-*.xlsx"))
        wb = openpyxl.load_workbook(path)
        self.assertEqual(
            wb.sheetnames,
            ["README", "Answers", "Questions", "Dimensions", "Overall"])
        dims = wb["Dimensions"]
        self.assertEqual(dims.max_row, 10)
        self.assertTrue(str(dims["D2"].value).startswith("=AVERAGE(")
                        or str(dims["D2"].value).startswith("=IF("))
        engine_d = {d["id"]: d["score"] for d in self.scores["dimensions"]}
        for row in dims.iter_rows(min_row=2, values_only=True):
            self.assertEqual(row[9], engine_d[row[0]])
        self.assertEqual(wb["Overall"]["C2"].value,
                         self.scores["overall"]["score"])

    def test_compare_same_round_has_zero_deltas(self) -> None:
        data = json.loads(MOCK.read_text())
        result = compare_rounds.compare_v2_v2(data, data)
        self.assertEqual(result["overall"]["delta"], 0.0)
        self.assertTrue(all(u["delta"] in (0.0, None)
                            for u in result["units"]))

    def test_compare_v1_to_v2_uses_lineage(self) -> None:
        before = json.loads((ROOT / "respostas.json.example").read_text())
        after = json.loads(MOCK.read_text())
        result = compare_rounds.compare_v1_v2(before, after)
        self.assertIsNone(result["overall"]["delta"])
        self.assertTrue(result["caveats"])
        fw = json.loads((ROOT / "framework.v2.json").read_text())
        with_lineage = sum(1 for d in fw["dimensions"]
                           for q in d["questions"] if q["v1_lineage"])
        self.assertEqual(len(result["questions"]), with_lineage)

    def test_report_payload_uses_engine_outputs(self) -> None:
        import build_report_v2
        payload = build_report_v2.build_payload(self.kit, self.out)
        self.assertEqual(payload["overall"], self.scores["overall"])
        self.assertEqual([g["id"] for g in payload["groups"]],
                         ["G1", "G2", "G3"])
        self.assertEqual(sum(len(g["dimensions"])
                             for g in payload["groups"]), 9)
        cited = {n for b in payload["backlog"] for n in b["basis"]}
        self.assertTrue(cited <= {r["n"] for r in payload["references"]})
        for d in payload["dimensions"]:
            for q in d["lowest"]:
                self.assertTrue(q["l3"])

    def test_implementation_guide_plan(self) -> None:
        import build_report_v2
        payload = build_report_v2.build_payload(self.kit, self.out)
        impl = payload["impl"]
        gaps = json.loads((self.out / "gaps.json").read_text())
        self.assertEqual(len(impl["with_gap"]), len(gaps["gaps"]))
        placed = [d["id"] for ph in impl["phases"] for d in ph["dimensions"]]
        self.assertEqual(sorted(placed),
                         sorted(g["dimension_id"] for g in gaps["gaps"]))
        self.assertFalse(impl["wizard_present"])
        self.assertTrue(all(v is None for v in impl["wizard"].values()))
        kinds = [r["kind"] for r in impl["risks"]]
        self.assertIn("amplification", kinds)
        (self.kit / "implementation-guide-inputs.json").write_text(
            json.dumps({"metadata": {"completion_pct": 20},
                        "implementation_guide_inputs": {
                            "dimension_owners": "D8: Ana, platform lead",
                            "raci_matrix": "(fill in manually)"}}),
            encoding="utf-8")
        impl = build_report_v2.build_payload(self.kit, self.out)["impl"]
        self.assertEqual(impl["owners"], {"D8": "Ana, platform lead"})
        self.assertIsNone(impl["wizard"]["raci_matrix"])
        self.assertEqual(impl["wizard_completion"], 20)

    def test_calculator_matches_engine(self) -> None:
        node = shutil.which("node")
        if not node:
            self.skipTest("node is not installed")
        import subprocess

        import generate_v2_tools_html as tools
        fw = json.loads((ROOT / "framework.v2.json").read_text("utf-8"))
        slim = tools.slim_framework(fw, True)
        inputs = {d["id"]: {"score": d["score"]}
                  for d in self.scores["dimensions"]}
        script = (tools.CALC_JS + "\nconsole.log(JSON.stringify(calc("
                  + json.dumps(slim) + ", " + json.dumps(inputs) + ")));")
        run = subprocess.run([node, "-e", script], capture_output=True,
                             text=True, check=True)
        got = json.loads(run.stdout)
        self.assertAlmostEqual(got["overall"],
                               self.scores["overall"]["score"], places=3)
        gaps = json.loads((self.out / "gaps.json").read_text())
        want = {g["dimension_id"]: g["priority"].split()[0]
                for g in gaps["gaps"]}
        have = {d["id"]: f"P{d['priority']}" for d in got["dims"]
                if d["priority"] is not None}
        self.assertEqual(have, want)
        recs = json.loads((self.out / "recomendacoes.json").read_text())
        self.assertEqual(
            [s["id"] for s in got["strategies"] if s["recommended"]],
            [s["strategy_id"] for s in recs["ranked_strategies"]])
        self.assertEqual(
            got["amplification"],
            [a["dimension_id"]
             for a in self.scores["flags"]["amplification_risk"]])

    def test_merge_offline_exports(self) -> None:
        import merge_offline_respostas as merge
        mock = json.loads(MOCK.read_text("utf-8"))
        folder = self.tmp / "exports"
        folder.mkdir()
        for idx, person in enumerate(mock["respondents"][:3]):
            one = {"metadata": {**mock["metadata"], "source":
                                "offline-html"},
                   "respondents": [{**person, "id": "R01",
                                    "name": "Respondent 01"}]}
            (folder / f"r{idx}.json").write_text(json.dumps(one),
                                                 encoding="utf-8")
        merged, _ = merge.merge(merge.input_files([str(folder)]), None,
                                None, False)
        ids = [p["id"] for p in merged["respondents"]]
        self.assertEqual(ids, ["R01", "R02", "R03"])
        self.assertEqual(len({p["name"] for p in merged["respondents"]}), 3)
        v1 = folder / "v1.json"
        v1.write_text(json.dumps({"metadata": {}, "responses": {}}))
        with self.assertRaises(ValueError):
            merge.merge([v1], None, None, False)

    def test_telemetry_and_repo_crosschecks(self) -> None:
        import build_report_v2
        import import_copilot_metrics as tel
        import scan_repos_ai_config as scan
        day = {"day": "2026-09-20", "monthly_active_users": 120,
               "totals_by_ai_adoption_phase": [
                   {"phase": "No Cohort", "users_in_phase_28d": 40},
                   {"phase": "Phase 1", "users_in_phase_28d": 70},
                   {"phase": "Phase 2", "users_in_phase_28d": 25},
                   {"phase": "Phase 3", "users_in_phase_28d": 5}]}
        report = self.tmp / "org-28-day.json"
        report.write_text(json.dumps({"report_end_day": "2026-09-20",
                                      "day_totals": [day]}))
        summary = tel.summarize(tel.read_records(report), 200)
        self.assertEqual(summary["population"], 140)
        self.assertEqual(summary["shares"]["monthly_active_of_seats"], 0.6)
        (self.out / "telemetria.json").write_text(json.dumps(summary))
        self.assertEqual(scan.classify(".github/copilot-instructions.md"),
                         ("rules", 2))
        self.assertEqual(scan.classify(".github/agents/x.agent.md"),
                         ("agents", 3))
        self.assertIsNone(scan.classify("src/main.py"))
        repos = [scan.scan_file_list("a", []),
                 scan.scan_file_list("b", ["AGENTS.md"]),
                 scan.scan_file_list("c", [".github/prompts/p.prompt.md"])]
        (self.out / "repo-scan.json").write_text(json.dumps(
            {"metadata": {"source": "test"},
             "summary": scan.summarize(repos), "repositories": repos}))
        payload = build_report_v2.build_payload(self.kit, self.out)
        checks = payload["checks"]
        self.assertEqual(checks["repo"]["d4q4"]["implied_level"], "L3")
        self.assertEqual(checks["telemetry"]["d4q1"]["implied_level"], "L2")
        self.assertTrue(checks["telemetry"]["d4q1"]["flag"])
        kinds = [r["kind"] for r in payload["impl"]["risks"]]
        self.assertIn("telemetry_gap", kinds)
        self.assertIn(47, {r["n"] for r in payload["references"]})

    def test_repo_scan_github_api_paths(self) -> None:
        import scan_repos_ai_config as scan
        calls = []

        def fake_api(url: str, token: str):
            calls.append(url)
            if "/orgs/" in url:
                return [] if "page=2" in url else [
                    {"name": "svc", "default_branch": "main"},
                    {"name": "old", "archived": True}]
            return {"tree": [{"type": "blob", "path": "AGENTS.md"},
                             {"type": "blob",
                              "path": ".github/agents/a.agent.md"}],
                    "truncated": False}
        original = scan._api
        scan._api = fake_api
        try:
            repos = scan.scan_github("contoso", "t", 10, False)
        finally:
            scan._api = original
        self.assertEqual([r["name"] for r in repos], ["svc"])
        self.assertEqual(repos[0]["level"], "L3")
        self.assertTrue(any("git/trees/main?recursive=1" in c
                            for c in calls))

    def test_compare_pdf(self) -> None:
        result = compare_rounds.compare_v2_v2(
            json.loads(MOCK.read_text("utf-8")),
            json.loads(MOCK.read_text("utf-8")))
        pdf = compare_rounds.render_pdf(
            result, json.loads(MOCK.read_text("utf-8")), self.out, "es")
        self.assertGreater(pdf.stat().st_size, 10_000)


if __name__ == "__main__":
    unittest.main()

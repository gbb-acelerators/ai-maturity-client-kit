#!/usr/bin/env python3
"""Generate the v2 collection artifacts from framework.v2.json.

Outputs (never edit them by hand, rerun this script):
- coleta/perguntas-para-forms.md      (PT-BR, canonical kit language)
- coleta/perguntas-para-forms.en.md   (EN)
- coleta/perguntas-para-forms.es.md   (ES)
- coleta/template-export-forms.xlsx   (Forms export shape, v2 columns)
- formularios/assessment-v2.html      (offline form, 3 languages,
                                        exports respostas.json)

Usage:
    python3 scripts/generate_v2_collection.py          # write files
    python3 scripts/generate_v2_collection.py --check  # fail if stale
"""
from __future__ import annotations

import argparse
import html
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BANKS = {"pt-br": "coleta/perguntas-para-forms.md",
         "en": "coleta/perguntas-para-forms.en.md",
         "es": "coleta/perguntas-para-forms.es.md"}
HTML_OUT = "formularios/assessment-v2.html"
XLSX_OUT = "coleta/template-export-forms.xlsx"

UI = {
    "en": {
        "title": "Microsoft Forms question bank: AI-Assisted SDLC "
                 "Maturity Assessment v2",
        "generated": "Generated from `framework.v2.json` (version "
                     "{version}) by `scripts/generate_v2_collection.py`. "
                     "Do not edit by hand. Source of the wording: "
                     "[AI-Maturity-Form-Questions_v2.md]"
                     "(AI-Maturity-Form-Questions_v2.md). The v1 bank "
                     "(158 questions) is archived in [v1/](v1/).",
        "how": "How to build the form",
        "steps": [
            "Go to <https://forms.office.com> and create a blank form. "
            "Suggested title: `AI-Assisted SDLC Maturity Assessment v2 - "
            "<Organization>`.",
            "Paste the privacy notice from [INSTRUCOES-FORMS.md]"
            "(INSTRUCOES-FORMS.md) into the form description.",
            "Add **10 sections**: Section 0 (profile) and one per "
            "dimension, D1 to D9.",
            "Section 0: add the 5 profile questions as **Choice**. `R-Q3` "
            "allows multiple answers. They are not scored.",
            "For each scored question add 2 elements: a **Choice** "
            "(single answer) whose title starts with the ID and a colon "
            "(for example `D4-Q3: ...`), with the 6 options below in "
            "order; and an optional **Long Text** titled "
            "`Evidence (<ID>)`.",
            "Paste the **L3 looks like** and **L4 looks like** lines into "
            "the question subtitle.",
            "Share the link. Aim for at least 3 respondents per role "
            "(`R-Q1`).",
            "`Responses` > `Open in Excel`, download the file, then run "
            "`make import XLSX=<file>`.",
        ],
        "count": "Element count: 5 profile + 61 scored + 61 optional "
                 "evidence fields = 127 elements.",
        "options": "The 6 options for every scored question",
        "prefix_warn": "Keep the `L0` to `L4` and `NA` prefix at the "
                       "start of each option and the ID at the start of "
                       "each title: the importer relies on both.",
        "rules": "How to answer",
        "rule_items": [
            "Read every question as \"to what extent is this true?\".",
            "Answer at the highest level where **every** part of the "
            "question is true.",
            "If coverage, governance and measurement point to different "
            "levels, pick the lowest.",
            "Choose `L0` when the practice could apply but does not "
            "exist yet; choose `NA` only when you do not know or the "
            "activity does not exist in your scope.",
        ],
        "section0": "Section 0: Respondent profile (not scored)",
        "multi": "Choice, multiple answers",
        "single": "Choice, single answer",
        "section": "Section {id}: {name}",
        "questions_n": "{n} questions. Why it matters: {why}",
        "unit": "Coverage unit",
        "units": {"teams": "teams", "engineers": "engineers",
                  "repositories": "repositories", "services": "services",
                  "organization": "organization-wide practice (use the "
                                  "governance and measurement columns)"},
        "l1_l2": "L1 to L2 look like",
        "l3": "L3 looks like",
        "l4": "L4 looks like",
        "evidence": "Evidence examples",
        "scope": "Scope note",
        "evidence_field": "Evidence field",
        "placeholder": "Tool, % coverage, metric, time window, link",
        "html_title": "AI-Assisted SDLC Maturity Assessment v2",
        "html_intro": "Offline form. Answers stay in this browser. Use "
                      "Export to download a respostas.json for the kit "
                      "(one respondent per file).",
        "org": "Organization",
        "name": "Your name (optional)",
        "export_json": "Export respostas.json",
        "export_csv": "Export CSV",
        "progress": "answered",
        "privacy": "Privacy: this page sends nothing over the network. "
                   "The exported file contains only what you type here.",
    },
    "pt-br": {
        "title": "Banco de perguntas para Microsoft Forms: AI-Assisted "
                 "SDLC Maturity Assessment v2",
        "generated": "Gerado a partir de `framework.v2.json` (versão "
                     "{version}) por `scripts/generate_v2_collection.py`. "
                     "Não edite à mão. Fonte do texto: "
                     "[AI-Maturity-Form-Questions_v2.md]"
                     "(AI-Maturity-Form-Questions_v2.md). O banco v1 (158 "
                     "perguntas) está arquivado em [v1/](v1/).",
        "how": "Como montar o formulário",
        "steps": [
            "Acesse <https://forms.office.com> e crie um formulário em "
            "branco. Título sugerido: `AI-Assisted SDLC Maturity "
            "Assessment v2 - <Organização>`.",
            "Cole o aviso de privacidade de [INSTRUCOES-FORMS.pt-br.md]"
            "(INSTRUCOES-FORMS.pt-br.md) na descrição do formulário.",
            "Adicione **10 seções**: Seção 0 (perfil) e uma por dimensão, "
            "D1 a D9.",
            "Seção 0: adicione as 5 perguntas de perfil como **Choice**. "
            "`R-Q3` permite várias respostas. Elas não pontuam.",
            "Para cada pergunta pontuada, adicione 2 elementos: um "
            "**Choice** (resposta única) cujo título começa com o ID e "
            "dois-pontos (por exemplo `D4-Q3: ...`), com as 6 opções "
            "abaixo na ordem; e um **Long Text** opcional com o título "
            "`Evidence (<ID>)`.",
            "Cole as linhas **L3 se parece com** e **L4 se parece com** "
            "no subtítulo da pergunta.",
            "Compartilhe o link. Busque ao menos 3 respondentes por papel "
            "(`R-Q1`).",
            "`Responses` > `Open in Excel`, baixe o arquivo e rode "
            "`make import XLSX=<arquivo>`.",
        ],
        "count": "Total de elementos: 5 de perfil + 61 pontuados + 61 "
                 "campos opcionais de evidência = 127 elementos.",
        "options": "As 6 opções de toda pergunta pontuada",
        "prefix_warn": "Mantenha o prefixo `L0` a `L4` e `NA` no início "
                       "de cada opção e o ID no início de cada título: o "
                       "importador depende dos dois. Mantenha também o "
                       "rótulo `Evidence (<ID>)` em inglês.",
        "rules": "Como responder",
        "rule_items": [
            "Leia cada pergunta como \"em que medida isto é verdade?\".",
            "Responda no nível mais alto em que **todas** as partes da "
            "pergunta são verdadeiras.",
            "Se cobertura, governança e medição apontarem níveis "
            "diferentes, escolha o menor.",
            "Escolha `L0` quando a prática poderia existir mas ainda não "
            "existe; escolha `NA` só quando você não sabe ou a atividade "
            "não existe no seu escopo.",
        ],
        "section0": "Seção 0: Perfil do respondente (não pontua)",
        "multi": "Choice, várias respostas",
        "single": "Choice, resposta única",
        "section": "Seção {id}: {name}",
        "questions_n": "{n} perguntas. Por que importa: {why}",
        "unit": "Unidade de cobertura",
        "units": {"teams": "times", "engineers": "engenheiros",
                  "repositories": "repositórios", "services": "serviços",
                  "organization": "prática da organização inteira (use "
                                  "as colunas de governança e medição)"},
        "l1_l2": "L1 a L2 se parecem com",
        "l3": "L3 se parece com",
        "l4": "L4 se parece com",
        "evidence": "Exemplos de evidência",
        "scope": "Nota de escopo",
        "evidence_field": "Campo de evidência",
        "placeholder": "Ferramenta, % de cobertura, métrica, período, "
                       "link",
        "html_title": "AI-Assisted SDLC Maturity Assessment v2",
        "html_intro": "Formulário offline. As respostas ficam neste "
                      "navegador. Use Exportar para baixar um "
                      "respostas.json para o kit (um respondente por "
                      "arquivo).",
        "org": "Organização",
        "name": "Seu nome (opcional)",
        "export_json": "Exportar respostas.json",
        "export_csv": "Exportar CSV",
        "progress": "respondidas",
        "privacy": "Privacidade: esta página não envia nada pela rede. O "
                   "arquivo exportado contém apenas o que você digitar "
                   "aqui.",
    },
    "es": {
        "title": "Banco de preguntas para Microsoft Forms: AI-Assisted "
                 "SDLC Maturity Assessment v2",
        "generated": "Generado a partir de `framework.v2.json` (versión "
                     "{version}) por `scripts/generate_v2_collection.py`. "
                     "No lo edite a mano. Fuente del texto: "
                     "[AI-Maturity-Form-Questions_v2.md]"
                     "(AI-Maturity-Form-Questions_v2.md). El banco v1 "
                     "(158 preguntas) está archivado en [v1/](v1/).",
        "how": "Cómo armar el formulario",
        "steps": [
            "Vaya a <https://forms.office.com> y cree un formulario en "
            "blanco. Título sugerido: `AI-Assisted SDLC Maturity "
            "Assessment v2 - <Organización>`.",
            "Pegue el aviso de privacidad de "
            "[../kit-es/INSTRUCCIONES-FORMS.md]"
            "(../kit-es/INSTRUCCIONES-FORMS.md) en la descripción del "
            "formulario.",
            "Agregue **10 secciones**: Sección 0 (perfil) y una por "
            "dimensión, D1 a D9.",
            "Sección 0: agregue las 5 preguntas de perfil como "
            "**Choice**. `R-Q3` permite varias respuestas. No puntúan.",
            "Para cada pregunta puntuada agregue 2 elementos: un "
            "**Choice** (respuesta única) cuyo título empieza con el ID y "
            "dos puntos (por ejemplo `D4-Q3: ...`), con las 6 opciones de "
            "abajo en orden; y un **Long Text** opcional con el título "
            "`Evidence (<ID>)`.",
            "Pegue las líneas **L3 se ve así** y **L4 se ve así** en el "
            "subtítulo de la pregunta.",
            "Comparta el enlace. Busque al menos 3 personas por rol "
            "(`R-Q1`).",
            "`Responses` > `Open in Excel`, descargue el archivo y "
            "ejecute `make import XLSX=<archivo>`.",
        ],
        "count": "Total de elementos: 5 de perfil + 61 puntuadas + 61 "
                 "campos opcionales de evidencia = 127 elementos.",
        "options": "Las 6 opciones de toda pregunta puntuada",
        "prefix_warn": "Mantenga el prefijo `L0` a `L4` y `NA` al inicio "
                       "de cada opción y el ID al inicio de cada título: "
                       "el importador depende de ambos. Mantenga también "
                       "la etiqueta `Evidence (<ID>)` en inglés.",
        "rules": "Cómo responder",
        "rule_items": [
            "Lea cada pregunta como \"¿en qué medida esto es cierto?\".",
            "Responda en el nivel más alto en que **todas** las partes de "
            "la pregunta son ciertas.",
            "Si cobertura, gobernanza y medición indican niveles "
            "distintos, elija el menor.",
            "Elija `L0` cuando la práctica podría aplicar pero todavía no "
            "existe; elija `NA` solo cuando no sabe o la actividad no "
            "existe en su alcance.",
        ],
        "section0": "Sección 0: Perfil de la persona que responde (no "
                    "puntúa)",
        "multi": "Choice, varias respuestas",
        "single": "Choice, respuesta única",
        "section": "Sección {id}: {name}",
        "questions_n": "{n} preguntas. Por qué importa: {why}",
        "unit": "Unidad de cobertura",
        "units": {"teams": "equipos", "engineers": "ingenieros",
                  "repositories": "repositorios", "services": "servicios",
                  "organization": "práctica de toda la organización (use "
                                  "las columnas de gobernanza y medición)"},
        "l1_l2": "L1 a L2 se ven así",
        "l3": "L3 se ve así",
        "l4": "L4 se ve así",
        "evidence": "Ejemplos de evidencia",
        "scope": "Nota de alcance",
        "evidence_field": "Campo de evidencia",
        "placeholder": "Herramienta, % de cobertura, métrica, período, "
                       "enlace",
        "html_title": "AI-Assisted SDLC Maturity Assessment v2",
        "html_intro": "Formulario offline. Las respuestas quedan en este "
                      "navegador. Use Exportar para descargar un "
                      "respostas.json para el kit (una persona por "
                      "archivo).",
        "org": "Organización",
        "name": "Su nombre (opcional)",
        "export_json": "Exportar respostas.json",
        "export_csv": "Exportar CSV",
        "progress": "respondidas",
        "privacy": "Privacidad: esta página no envía nada por la red. El "
                   "archivo exportado contiene solo lo que usted escriba "
                   "aquí.",
    },
}


def load_fw() -> dict:
    return json.loads((ROOT / "framework.v2.json").read_text("utf-8"))


def bank_md(fw: dict, lang: str) -> str:
    t = UI[lang]
    out = [f"# {t['title']}", "",
           f"> {t['generated'].format(version=fw['version'])}", "",
           f"## {t['how']}", ""]
    out += [f"{i}. {s}" for i, s in enumerate(t["steps"], start=1)]
    out += ["", t["count"], "", f"## {t['options']}", ""]
    out += [f"- **{o}**" for o in fw["options"][lang]]
    out += ["", f"> {t['prefix_warn']}", "", f"## {t['rules']}", ""]
    out += [f"- {r}" for r in t["rule_items"]]
    out += ["", "---", "", f"## {t['section0']}", ""]
    for p in fw["profile_questions"]:
        kind = t["multi"] if p["multi"] else t["single"]
        out += [f"### {p['id']}: {p['title'][lang]}", "",
                f"**{p['id']}: {p['text'][lang]}**", "", f"_{kind}_", ""]
        out += [f"- {o}" for o in p["options"][lang]]
        out.append("")
    for d in fw["dimensions"]:
        out += ["---", "",
                "## " + t["section"].format(id=d["id"],
                                            name=d["name"][lang]), "",
                "_" + t["questions_n"].format(
                    n=len(d["questions"]),
                    why=d["why_it_matters"][lang]) + "_", ""]
        for q in d["questions"]:
            out += [f"### {q['id']}: {q['title'][lang]}", "",
                    f"**{q['id']}: {q['text'][lang]}**", ""]
            if "scope_note" in q:
                out.append(f"- **{t['scope']}:** {q['scope_note'][lang]}")
            out.append(f"- **{t['unit']}:** {t['units'][q['unit']]}")
            a = q["anchors"]
            if "l1_l2" in a:
                out.append(f"- **{t['l1_l2']}:** {a['l1_l2'][lang]}")
            out += [f"- **{t['l3']}:** {a['l3'][lang]}",
                    f"- **{t['l4']}:** {a['l4'][lang]}",
                    f"- **{t['evidence']}:** "
                    f"{q['evidence_examples'][lang]}",
                    f"- **{t['evidence_field']}:** `Evidence ({q['id']})`"
                    f" · _{t['placeholder']}_", ""]
    return "\n".join(out).rstrip() + "\n"


MS_LOGO = (
    '<svg viewBox="0 0 23 23" width="18" height="18" aria-label="Microsoft"'
    ' role="img"><rect x="1" y="1" width="10" height="10" fill="#F25022"/>'
    '<rect x="12" y="1" width="10" height="10" fill="#7FBA00"/>'
    '<rect x="1" y="12" width="10" height="10" fill="#00A4EF"/>'
    '<rect x="12" y="12" width="10" height="10" fill="#FFB900"/></svg>')

HTML_TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>AI-Assisted SDLC Maturity Assessment v2</title>
<!-- Generated by scripts/generate_v2_collection.py. Do not edit. -->
<style>
:root{--primary:#0078D4;--ink:#1b1b1f;--muted:#5b5b66;--line:#e3e3e8;
--bg:#f7f7f9;--l0:#9aa0a6;--l1:#f7b500;--l2:#f08a24;--l3:#2b88d8;
--l4:#107c10}
*{box-sizing:border-box}body{margin:0;font-family:"Segoe UI",system-ui,
sans-serif;color:var(--ink);background:var(--bg);line-height:1.45}
header{position:sticky;top:0;background:#fff;border-bottom:1px solid
var(--line);padding:10px 16px;z-index:5}
.bar{display:flex;gap:12px;align-items:center;flex-wrap:wrap;
max-width:960px;margin:0 auto}
.brand{display:flex;gap:8px;align-items:center;font-size:13px;
color:var(--muted)}.grow{flex:1}
main{max-width:960px;margin:0 auto;padding:16px}
section{background:#fff;border:1px solid var(--line);border-radius:8px;
padding:16px;margin:16px 0}
h1{font-size:22px;margin:8px 0}h2{font-size:18px;margin:0 0 4px}
.why{color:var(--muted);font-size:13px;margin:0 0 12px}
.q{border-top:1px solid var(--line);padding:12px 0}
.q:first-of-type{border-top:0}.qid{font-weight:700;color:var(--primary)}
.anc{font-size:13px;color:var(--muted);margin:6px 0}
.opts label{display:block;padding:4px 6px;border-radius:4px;
cursor:pointer}.opts label:hover{background:var(--bg)}
textarea{width:100%;min-height:44px;font:inherit;border:1px solid
var(--line);border-radius:4px;padding:6px}
input[type=text]{font:inherit;padding:6px;border:1px solid var(--line);
border-radius:4px;min-width:220px}
button,select{font:inherit;padding:6px 12px;border-radius:4px;border:1px
solid var(--primary);background:#fff;color:var(--primary);cursor:pointer}
button.primary{background:var(--primary);color:#fff}
.progress{font-size:13px;color:var(--muted)}
.note{font-size:12px;color:var(--muted)}
footer{max-width:960px;margin:0 auto;padding:16px;font-size:12px;
color:var(--muted)}
</style>
</head>
<body>
<header><div class="bar">
<span class="brand">__LOGO__<span>Paula Silva | Global Developer Solutions
Advisor</span></span><span class="grow"></span>
<span class="progress" id="progress"></span>
<select id="lang" aria-label="Language"><option value="en">English</option>
<option value="pt-br">Português (Brasil)</option>
<option value="es">Español</option></select>
<button id="csv"></button><button id="json" class="primary"></button>
</div></header>
<main>
<h1 id="title"></h1><p id="intro"></p><p class="note" id="privacy"></p>
<p><label><span id="orgLabel"></span><br><input type="text" id="org">
</label> <label><span id="nameLabel"></span><br>
<input type="text" id="name"></label></p>
<div id="form"></div>
</main>
<footer>framework v__VERSION__ · coleta/AI-Maturity-Form-Questions_v2.md
</footer>
<script>
const FW = __DATA__;
const UI = __UI__;
const KEY = "ai-maturity-v2-" + FW.version;
let state = JSON.parse(localStorage.getItem(KEY) || '{"answers":{},'
  + '"profile":{}}');
let lang = localStorage.getItem(KEY + "-lang") || "en";
const $ = (id) => document.getElementById(id);
const esc = (s) => String(s).replace(/[&<>"]/g, (c) => ({"&":"&amp;",
  "<":"&lt;", ">":"&gt;", '"':"&quot;"}[c]));
function save() { localStorage.setItem(KEY, JSON.stringify(state));
  progress(); }
function progress() {
  const total = FW.dimensions.reduce((n, d) => n + d.questions.length, 0);
  const done = Object.values(state.answers).filter((a) => a.level !==
    undefined).length;
  $("progress").textContent = done + " / " + total + " " + UI[lang].progress;
}
function render() {
  const t = UI[lang];
  document.documentElement.lang = lang;
  $("title").textContent = t.html_title; $("intro").textContent =
    t.html_intro; $("privacy").textContent = t.privacy;
  $("orgLabel").textContent = t.org; $("nameLabel").textContent = t.name;
  $("json").textContent = t.export_json; $("csv").textContent =
    t.export_csv;
  $("org").value = state.org || ""; $("name").value = state.name || "";
  let h = '<section><h2>' + esc(t.section0) + '</h2>';
  for (const p of FW.profile_questions) {
    h += '<div class="q"><div><span class="qid">' + p.id + '</span> '
      + esc(p.text[lang]) + '</div><div class="opts">';
    p.options[lang].forEach((o, i) => {
      const canon = p.options.en[i];
      const cur = state.profile[p.id];
      const on = p.multi ? (cur || []).includes(canon) : cur === canon;
      h += '<label><input type="' + (p.multi ? "checkbox" : "radio")
        + '" name="' + p.id + '" data-p="' + p.id + '" value="'
        + esc(canon) + '"' + (on ? " checked" : "") + '> ' + esc(o)
        + '</label>';
    });
    h += '</div></div>';
  }
  h += '</section>';
  for (const d of FW.dimensions) {
    h += '<section><h2>' + d.id + ' · ' + esc(d.name[lang]) + '</h2>'
      + '<p class="why">' + esc(d.why_it_matters[lang]) + '</p>';
    for (const q of d.questions) {
      const a = state.answers[q.id] || {};
      h += '<div class="q"><div><span class="qid">' + q.id + '</span> '
        + esc(q.text[lang]) + '</div>';
      if (q.anchors.l1_l2) h += '<div class="anc"><b>' + esc(t.l1_l2)
        + ':</b> ' + esc(q.anchors.l1_l2[lang]) + '</div>';
      h += '<div class="anc"><b>' + esc(t.l3) + ':</b> '
        + esc(q.anchors.l3[lang]) + '</div><div class="anc"><b>'
        + esc(t.l4) + ':</b> ' + esc(q.anchors.l4[lang]) + '</div>'
        + '<div class="opts">';
      FW.options[lang].forEach((o, i) => {
        const val = i === 5 ? "na" : String(i);
        const on = a.level === undefined ? false :
          (a.level === null ? val === "na" : String(a.level) === val);
        h += '<label><input type="radio" name="' + q.id + '" data-q="'
          + q.id + '" value="' + val + '"' + (on ? " checked" : "")
          + '> ' + esc(o) + '</label>';
      });
      h += '</div><textarea data-e="' + q.id + '" placeholder="'
        + esc(t.placeholder) + '">' + esc(a.evidence || "")
        + '</textarea></div>';
    }
    h += '</section>';
  }
  $("form").innerHTML = h; progress();
}
document.addEventListener("change", (ev) => {
  const el = ev.target;
  if (el.dataset.q) {
    const a = state.answers[el.dataset.q] || {evidence: ""};
    a.level = el.value === "na" ? null : Number(el.value);
    state.answers[el.dataset.q] = a;
  } else if (el.dataset.p) {
    const p = FW.profile_questions.find((x) => x.id === el.dataset.p);
    if (p.multi) {
      const cur = new Set(state.profile[p.id] || []);
      el.checked ? cur.add(el.value) : cur.delete(el.value);
      state.profile[p.id] = [...cur];
    } else { state.profile[p.id] = el.value; }
  } else if (el.id === "lang") {
    lang = el.value; localStorage.setItem(KEY + "-lang", lang); render();
    return;
  }
  save();
});
document.addEventListener("input", (ev) => {
  const el = ev.target;
  if (el.dataset.e) {
    const a = state.answers[el.dataset.e] || {};
    a.evidence = el.value; state.answers[el.dataset.e] = a;
  } else if (el.id === "org") { state.org = el.value; }
  else if (el.id === "name") { state.name = el.value; }
  save();
});
function download(name, text, type) {
  const url = URL.createObjectURL(new Blob([text], {type}));
  const a = document.createElement("a"); a.href = url; a.download = name;
  a.click(); URL.revokeObjectURL(url);
}
$("json").onclick = () => {
  const answers = {};
  for (const [k, v] of Object.entries(state.answers)) {
    if (v.level !== undefined) answers[k] = {level: v.level,
      evidence: v.evidence || ""};
  }
  const out = {metadata: {organization: state.org || null,
    assessment_date: new Date().toISOString().slice(0, 10),
    language: lang === "pt-br" ? "pt-BR" : lang, source: "offline-html",
    framework_version: FW.version}, target_overrides: {},
    dimension_weights: {}, respondents: [{id: "R01",
    name: state.name || "Respondent 01", email: "", profile: state.profile,
    answers}]};
  download("respostas.json", JSON.stringify(out, null, 2),
    "application/json");
};
$("csv").onclick = () => {
  const rows = [["id", "level", "evidence"]];
  for (const d of FW.dimensions) for (const q of d.questions) {
    const a = state.answers[q.id] || {};
    rows.push([q.id, a.level === undefined ? "" : (a.level === null ?
      "NA" : "L" + a.level), a.evidence || ""]);
  }
  download("answers.csv", rows.map((r) => r.map((c) => '"' +
    String(c).replace(/"/g, '""') + '"').join(",")).join("\\n"),
    "text/csv");
};
$("lang").value = lang;
render();
</script>
</body>
</html>
"""


def html_form(fw: dict) -> str:
    slim = {
        "version": fw["version"],
        "options": fw["options"],
        "profile_questions": [
            {k: p[k] for k in ("id", "text", "multi", "options")}
            for p in fw["profile_questions"]],
        "dimensions": [
            {"id": d["id"], "name": d["name"],
             "why_it_matters": d["why_it_matters"],
             "questions": [{"id": q["id"], "text": q["text"],
                            "anchors": q["anchors"]}
                           for q in d["questions"]]}
            for d in fw["dimensions"]],
    }
    ui = {lang: {k: UI[lang][k] for k in (
        "html_title", "html_intro", "privacy", "org", "name",
        "export_json", "export_csv", "progress", "section0", "l1_l2",
        "l3", "l4", "placeholder")} for lang in UI}
    data = json.dumps(slim, ensure_ascii=False).replace("</", "<\\/")
    return (HTML_TEMPLATE.replace("__DATA__", data)
            .replace("__UI__", json.dumps(ui, ensure_ascii=False))
            .replace("__LOGO__", MS_LOGO)
            .replace("__VERSION__", html.escape(fw["version"])))


def write_xlsx(fw: dict, path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    header = ["ID", "Start time", "Completion time", "Email", "Name"]
    header += [f"{p['id']}: {p['text']['en']}"
               for p in fw["profile_questions"]]
    for d in fw["dimensions"]:
        for q in d["questions"]:
            header += [f"{q['id']}: {q['text']['en']}",
                       f"Evidence ({q['id']})"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    info = wb.create_sheet("README")
    for line in (
        "Microsoft Forms export template for framework v2 "
        f"({fw['version']}).",
        "One row per respondent. Keep the header row unchanged: the "
        "importer maps columns by the ID prefix.",
        "Scored columns: L0 to L4 or NA (prefix at the start). R-Q3 "
        "accepts several options separated by ';'.",
        "An illustrative filled example is coleta/v2-mock-forms-export"
        ".xlsx (synthetic data).",
        "Generated by scripts/generate_v2_collection.py.",
    ):
        info.append([line])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def outputs(fw: dict) -> dict[str, str]:
    out = {path: bank_md(fw, lang) for lang, path in BANKS.items()}
    out[HTML_OUT] = html_form(fw)
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args()
    fw = load_fw()
    files = outputs(fw)
    if args.check:
        stale = [p for p, text in files.items()
                 if not (ROOT / p).exists()
                 or (ROOT / p).read_text("utf-8") != text]
        if stale:
            print(f"✗ stale generated files: {stale}; run "
                  f"scripts/generate_v2_collection.py", file=sys.stderr)
            return 1
        print(f"✓ {len(files)} generated collection files up to date")
        return 0
    for rel, text in files.items():
        (ROOT / rel).parent.mkdir(parents=True, exist_ok=True)
        (ROOT / rel).write_text(text, encoding="utf-8")
        print(f"✓ {rel}")
    write_xlsx(fw, ROOT / XLSX_OUT)
    print(f"✓ {XLSX_OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

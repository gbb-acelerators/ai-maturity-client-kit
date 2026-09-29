#!/usr/bin/env python3
"""Import a Microsoft Forms Excel export into responses.json.

Deterministic implementation of the /import-responses skill:
maps question columns by ID prefix, parses the L0-L4 / NA options,
averages levels across respondents (no rounding), concatenates evidence,
backs up the previous responses.json, and writes an import log to output/.

The framework version is detected from the column IDs: `P#-C#-Q#`
columns import as v1 (averaged responses), `D#-Q#` and `R-Q#` columns
import as v2 (one entry per respondent with the profile answers, which
the v2 engine needs for persona scores and flags).

Usage:
    python3 scripts/import_forms_excel.py [forms-responses.xlsx]
    python3 scripts/import_forms_excel.py FILE --organization "Contoso"
"""
from __future__ import annotations

import argparse
import copy
import datetime
import json
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
QID_RE = re.compile(r"^\s*([A-Z]\d*-(?:C\d+-)?Q\d+)\b")
EVIDENCE_RE = re.compile(r"^\s*Evid(?:ência|ence|encia)\s*\(([^)]+)\)",
                         re.IGNORECASE)
LEVEL_RE = re.compile(r"^\s*(L[0-4]|NA)\b", re.IGNORECASE)
NA_WORDS = {"não sei", "no sé", "no se", "i do not know", "i don't know",
            "n/a", "na"}
NAME_HEADERS = {"name", "nome", "nombre"}
EMAIL_HEADERS = {"email", "e-mail", "correo", "correo electrónico"}
MIN_COVERAGE = 0.6
V2_ID_RE = re.compile(r"^(?:D\d+|R)-Q\d+$")
MULTI_SEPARATOR = ";"

LOG_TEXT = {
    "en": {
        "title": "Import log: {date}",
        "summary": "Summary",
        "file": "Imported file",
        "respondents": "Respondents",
        "processed": "Questions found in the file",
        "answered": "Questions with at least 1 answer",
        "backup": "Backup of the previous responses.json",
        "none": "none (no previous file)",
        "coverage": "Coverage per respondent",
        "cols": "| Respondent | Email | Answered | Evidence |",
        "alerts": "Alerts",
        "no_alerts": "No alerts.",
        "next": "Next step",
        "next_text": "Run `python3 scripts/assessment_engine.py all` "
                     "(or `/full-pipeline`).",
        "unknown": "row {row} ({name}): unrecognized value at {qid}: "
                   "{value!r}, treated as null",
        "unanswered": "{n} question(s) with no answer from any "
                      "respondent stay null",
        "missing": "{n} framework question(s) not found in the file",
        "extra": "column ID(s) not in framework.json were ignored: {ids}",
        "org": "metadata.organization is empty: set it (or use "
               "--organization) before generating PDFs",
    },
    "pt-br": {
        "title": "Log de importação: {date}",
        "summary": "Resumo",
        "file": "Arquivo importado",
        "respondents": "Respondentes",
        "processed": "Perguntas encontradas no arquivo",
        "answered": "Perguntas com ao menos 1 resposta",
        "backup": "Backup do responses.json anterior",
        "none": "nenhum (não havia arquivo)",
        "coverage": "Cobertura por respondente",
        "cols": "| Respondente | E-mail | Respondidas | Evidências |",
        "alerts": "Alertas",
        "no_alerts": "Nenhum alerta.",
        "next": "Próximo passo",
        "next_text": "Rode `python3 scripts/assessment_engine.py all` "
                     "(ou `/full-pipeline`).",
        "unknown": "linha {row} ({name}): valor não reconhecido em {qid}: "
                   "{value!r}, tratado como null",
        "unanswered": "{n} pergunta(s) sem resposta de nenhum respondente "
                      "ficam null",
        "missing": "{n} pergunta(s) do framework não encontradas no "
                   "arquivo",
        "extra": "IDs de coluna fora do framework.json foram ignorados: "
                 "{ids}",
        "org": "metadata.organization está vazio: preencha (ou use "
               "--organization) antes de gerar os PDFs",
    },
}


class FormsImportError(ValueError):
    pass


def framework_qids(framework: dict) -> list[str]:
    return [
        q["id"]
        for p in framework["pillars"]
        for c in p["capabilities"]
        for q in c["questions"]
    ]


def parse_level(value) -> tuple[float | None, bool]:
    """Return (level, recognized). Empty and NA are recognized nulls."""
    if value is None or str(value).strip() == "":
        return None, True
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return (float(value), True) if 0 <= value <= 4 else (None, False)
    text = str(value).strip()
    match = LEVEL_RE.match(text)
    if match:
        token = match.group(1).upper()
        return (None, True) if token == "NA" else (float(token[1]), True)
    if text.lower() in NA_WORDS:
        return None, True
    return None, False


def find_sheet(workbook):
    for ws in workbook.worksheets:
        header = [str(c.value or "") for c in ws[1]]
        if any(QID_RE.match(h) for h in header):
            return ws
    raise FormsImportError(
        "No column header starts with a question ID such as "
        "'D1-Q1:' (v2) or 'P1-C1-Q1:' (v1). Is this a Microsoft Forms "
        "export of the kit form?")


def map_columns(ws, known: set[str]) -> dict:
    qcols, ecols, extra = {}, {}, set()
    name_col = email_col = None
    for idx, cell in enumerate(ws[1]):
        header = str(cell.value or "").strip()
        low = header.lower()
        if (m := EVIDENCE_RE.match(header)):
            qid = m.group(1).strip()
            if qid in known:
                ecols[qid] = idx
            continue
        if (m := QID_RE.match(header)):
            qid = m.group(1)
            if qid in known:
                qcols[qid] = idx
            else:
                extra.add(qid)
            continue
        if low in NAME_HEADERS and name_col is None:
            name_col = idx
        elif low in EMAIL_HEADERS and email_col is None:
            email_col = idx
    return {"q": qcols, "e": ecols, "extra": sorted(extra),
            "name": name_col, "email": email_col}


def read_respondents(ws, cols: dict) -> tuple[list[dict], list[dict]]:
    respondents, unknown = [], []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True),
                                 start=2):
        def cell(idx):
            return row[idx] if idx is not None and idx < len(row) else None
        name = str(cell(cols["name"]) or "").strip()
        email = str(cell(cols["email"]) or "").strip()
        answers = {}
        for qid, idx in cols["q"].items():
            level, ok = parse_level(cell(idx))
            if not ok:
                unknown.append({"row": row_no, "name": name, "qid": qid,
                                "value": cell(idx)})
            evidence = str(cell(cols["e"].get(qid)) or "").strip()
            explicit = str(cell(idx) or "").strip() != ""
            if level is not None or evidence or explicit:
                answers[qid] = {"level": level, "evidence": evidence}
        if not answers and not name and not email:
            continue
        respondents.append({
            "row": row,
            "row_no": row_no,
            "name": name or f"Respondent {len(respondents) + 1}",
            "email": email,
            "answers": answers,
        })
    return respondents, unknown


def aggregate(respondents: list[dict], qids: list[str]) -> dict:
    multi = len(respondents) > 1
    result = {}
    for qid in qids:
        levels, evidence = [], []
        for r in respondents:
            ans = r["answers"].get(qid)
            if not ans:
                continue
            if ans["level"] is not None:
                levels.append(ans["level"])
            if ans["evidence"]:
                evidence.append(f"[{r['name']}]: {ans['evidence']}"
                                if multi else ans["evidence"])
        result[qid] = {
            "level": sum(levels) / len(levels) if levels else None,
            "evidence": "\n".join(evidence),
            "n_respondents": len(levels),
        }
    return result


def base_responses(target: Path) -> dict:
    if target.exists():
        return json.loads(target.read_text(encoding="utf-8"))
    example = json.loads(
        (ROOT / "responses.json.example").read_text(encoding="utf-8"))
    data = copy.deepcopy(example)
    data["metadata"] = {"language": "en"}
    data["target_overrides"] = {}
    for entry in data["responses"].values():
        entry["level"] = None
        entry["evidence"] = ""
    return data


def log_locale(meta: dict, override: str | None) -> str:
    raw = (override or meta.get("language") or "en").lower()
    return "pt-br" if raw.startswith("pt") else "en"


def write_log(path: Path, t: dict, info: dict) -> None:
    lines = [f"# {t['title'].format(date=info['date'])}", "",
             f"## {t['summary']}", "",
             f"- {t['file']}: `{info['file']}`",
             f"- {t['respondents']}: {len(info['respondents'])} "
             f"({', '.join(r['name'] for r in info['respondents'])})",
             f"- {t['processed']}: {info['found']} / {info['total']}",
             f"- {t['answered']}: {info['answered']}",
             f"- {t['backup']}: {info['backup'] or t['none']}", "",
             f"## {t['coverage']}", "", t["cols"],
             "| --- | --- | --- | --- |"]
    for r in info["respondents"]:
        answered = sum(1 for a in r["answers"].values()
                       if a["level"] is not None)
        evid = sum(1 for a in r["answers"].values() if a["evidence"])
        lines.append(f"| {r['name']} | {r['email'] or '-'} | "
                     f"{answered} / {info['total']} | {evid} |")
    lines += ["", f"## {t['alerts']}", ""]
    alerts = [f"- {a}" for a in info["alerts"]]
    lines += alerts or [t["no_alerts"]]
    lines += ["", f"## {t['next']}", "", t["next_text"], ""]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def is_v2_sheet(ws) -> bool:
    ids = [m.group(1) for c in ws[1]
           if (m := QID_RE.match(str(c.value or "")))]
    return bool(ids) and sum(bool(V2_ID_RE.match(i)) for i in ids) \
        > len(ids) / 2


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", str(text).replace("\u2013", "-")
                  .replace("\u2014", "-")).strip().lower()


def profile_maps(fw: dict) -> dict[str, dict[str, str]]:
    """Map option text in any language to the canonical English text."""
    maps = {}
    for p in fw["profile_questions"]:
        canon = p["options"]["en"]
        table = {}
        for opts in p["options"].values():
            for idx, text in enumerate(opts):
                table[_norm(text)] = canon[idx]
        maps[p["id"]] = table
    return maps


def read_profile(value, qid: str, multi: bool, maps: dict,
                 unknown: list, row_no: int, name: str):
    if value is None or str(value).strip() == "":
        return [] if multi else None
    parts = [x for x in str(value).split(MULTI_SEPARATOR) if x.strip()] \
        if multi else [str(value)]
    out = []
    for part in parts:
        canon = maps[qid].get(_norm(part))
        if canon is None:
            unknown.append({"row": row_no, "name": name, "qid": qid,
                            "value": part.strip()})
            canon = part.strip()
        out.append(canon)
    return out if multi else out[0]


def run_v2(args, ws, xlsx: Path) -> int:
    fw = json.loads((ROOT / "framework.v2.json").read_text("utf-8"))
    qids = [q["id"] for d in fw["dimensions"] for q in d["questions"]]
    pids = [p["id"] for p in fw["profile_questions"]]
    multi = {p["id"]: p["multi"] for p in fw["profile_questions"]}
    cols = map_columns(ws, set(qids) | set(pids))
    pcols = {k: v for k, v in cols["q"].items() if k in pids}
    cols["q"] = {k: v for k, v in cols["q"].items() if k in qids}
    found = len(cols["q"])
    if found < MIN_COVERAGE * len(qids) and not args.allow_partial:
        raise FormsImportError(
            f"Only {found} of {len(qids)} v2 questions were found in the "
            f"file. Question titles must start with the ID (for example "
            f"'D4-Q3: ...'), or pass --allow-partial.")
    respondents, unknown = read_respondents(ws, cols)
    if not respondents:
        raise FormsImportError("The file has no respondent rows.")
    maps = profile_maps(fw)
    for idx, person in enumerate(respondents, start=1):
        row = person["row"]
        person["id"] = f"R{idx:02d}"
        person["profile"] = {
            pid: read_profile(row[col] if col < len(row) else None, pid,
                              multi[pid], maps, unknown, person["row_no"],
                              person["name"])
            for pid, col in pcols.items()
        }

    target = Path(args.responses)
    old = json.loads(target.read_text("utf-8")) if target.exists() else {}
    backup = None
    if target.exists():
        ts = datetime.datetime.now(datetime.timezone.utc).strftime(
            "%Y%m%dT%H%M%S")
        backup = target.with_name(f"{target.name}.backup-{ts}")
        shutil.copy2(target, backup)
    meta = old.get("metadata", {})
    org = args.organization or meta.get("organization")
    data = {
        "metadata": {
            "organization": org,
            "assessment_date": datetime.date.today().isoformat(),
            "language": meta.get("language", "en"),
            "source": "microsoft-forms-import",
            "framework_version": fw["version"],
        },
        "target_overrides": old.get("target_overrides", {})
        if str(meta.get("framework_version", "")).startswith("2") else {},
        "dimension_weights": old.get("dimension_weights", {}),
        "respondents": [
            {"id": r["id"], "name": r["name"], "email": r["email"],
             "profile": r["profile"], "answers": r["answers"]}
            for r in respondents
        ],
    }
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n",
                      encoding="utf-8")
    t = LOG_TEXT[log_locale(data["metadata"], args.lang)]
    alerts = [t["unknown"].format(**u) for u in unknown]
    answered = {q for r in respondents for q, a in r["answers"].items()
                if a["level"] is not None}
    if len(qids) - len(answered):
        alerts.append(t["unanswered"].format(n=len(qids) - len(answered)))
    if found < len(qids):
        alerts.append(t["missing"].format(n=len(qids) - found))
    if len(pcols) < len(pids):
        alerts.append(t["missing"].format(n=len(pids) - len(pcols))
                      + " (profile R-Q#)")
    if cols["extra"]:
        alerts.append(t["extra"].format(ids=", ".join(cols["extra"])))
    if not org:
        alerts.append(t["org"])
    date = datetime.date.today().isoformat()
    log_path = Path(args.log_dir) / f"import-log-{date}.md"
    write_log(log_path, t, {
        "date": date, "file": xlsx.name, "respondents": respondents,
        "found": found, "total": len(qids), "answered": len(answered),
        "backup": backup.name if backup else None, "alerts": alerts,
    })
    print(f"✓ {target.name} (v2 {fw['version']}): {len(respondents)} "
          f"respondent(s), {len(answered)}/{len(qids)} questions answered")
    if backup:
        print(f"✓ Backup: {backup.name}")
    print(f"✓ Log: {log_path}")
    if alerts:
        print(f"⚠️ {len(alerts)} alert(s), see the log")
    return 0


def run(args) -> int:
    try:
        import openpyxl
    except ImportError:
        print("✗ openpyxl is required: make install-deps", file=sys.stderr)
        return 1
    xlsx = Path(args.xlsx)
    if not xlsx.exists():
        raise FormsImportError(f"File not found: {xlsx}")
    ws = find_sheet(openpyxl.load_workbook(xlsx, read_only=False,
                                           data_only=True))
    if is_v2_sheet(ws):
        return run_v2(args, ws, xlsx)
    framework = json.loads((ROOT / "framework.json").read_text("utf-8"))
    qids = framework_qids(framework)
    cols = map_columns(ws, set(qids))
    found = len(cols["q"])
    if found < MIN_COVERAGE * len(qids) and not args.allow_partial:
        raise FormsImportError(
            f"Only {found} of {len(qids)} framework questions were found "
            f"in the file. Check the form (question titles must start with "
            f"the ID) or pass --allow-partial.")
    respondents, unknown = read_respondents(ws, cols)
    if not respondents:
        raise FormsImportError("The file has no respondent rows.")
    agg = aggregate(respondents, qids)

    target = Path(args.responses)
    data = base_responses(target)
    backup = None
    if target.exists():
        ts = datetime.datetime.now(datetime.timezone.utc).strftime(
            "%Y%m%dT%H%M%S")
        backup = target.with_name(f"{target.name}.backup-{ts}")
        shutil.copy2(target, backup)
    meta = data.get("metadata", {})
    org = args.organization or meta.get("organization")
    data["metadata"] = {
        **{k: v for k, v in meta.items()
           if k not in ("respondent_name", "respondent_email")},
        "respondent_name": (respondents[0]["name"] if len(respondents) == 1
                            else f"Aggregate of {len(respondents)} "
                            f"respondents"),
        "respondent_email": (respondents[0]["email"]
                             if len(respondents) == 1 else "-"),
        "respondent_role": (meta.get("respondent_role")
                            if len(respondents) == 1
                            else "Multi-respondent"),
        "organization": org,
        "assessment_date": datetime.date.today().isoformat(),
        "language": meta.get("language", "en"),
        "source": "microsoft-forms-import",
        "respondents": [{"name": r["name"], "email": r["email"]}
                        for r in respondents],
        "framework_version": framework.get("version"),
    }
    for qid, body in agg.items():
        entry = data["responses"].setdefault(qid, {})
        entry.update(body)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n",
                      encoding="utf-8")

    t = LOG_TEXT[log_locale(data["metadata"], args.lang)]
    alerts = [t["unknown"].format(**u) for u in unknown]
    unanswered = sum(1 for b in agg.values() if b["level"] is None)
    if unanswered:
        alerts.append(t["unanswered"].format(n=unanswered))
    if found < len(qids):
        alerts.append(t["missing"].format(n=len(qids) - found))
    if cols["extra"]:
        alerts.append(t["extra"].format(ids=", ".join(cols["extra"])))
    if not org:
        alerts.append(t["org"])
    date = datetime.date.today().isoformat()
    log_path = Path(args.log_dir) / f"import-log-{date}.md"
    write_log(log_path, t, {
        "date": date, "file": xlsx.name, "respondents": respondents,
        "found": found, "total": len(qids),
        "answered": len(qids) - unanswered,
        "backup": backup.name if backup else None, "alerts": alerts,
    })
    evidence = sum(1 for b in agg.values() if b["evidence"])
    print(f"✓ {target.name}: {len(respondents)} respondent(s), "
          f"{len(qids) - unanswered}/{len(qids)} questions answered, "
          f"{evidence} with evidence")
    if backup:
        print(f"✓ Backup: {backup.name}")
    print(f"✓ Log: {log_path}")
    if alerts:
        print(f"⚠️ {len(alerts)} alert(s), see the log")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("xlsx", nargs="?",
                    default=str(ROOT / "forms-responses.xlsx"))
    ap.add_argument("--responses", "--respostas", dest="responses",
                    default=str(ROOT / "responses.json"),
                    help="output file (default: responses.json)")
    ap.add_argument("--log-dir", default=str(ROOT / "output"))
    ap.add_argument("--organization", default=None)
    ap.add_argument("--lang", choices=("en", "pt-br"), default=None,
                    help="Log language (default: metadata.language)")
    ap.add_argument("--allow-partial", action="store_true",
                    help="Import even if fewer than 60%% of the "
                         "questions are in the file")
    args = ap.parse_args()
    try:
        return run(args)
    except FormsImportError as exc:
        print(f"✗ {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())

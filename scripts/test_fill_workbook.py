#!/usr/bin/env python3
"""Tests for scripts/fill_workbook.py (needs openpyxl).

If the optional `formulas` package is installed, the workbook formulas
are also evaluated and compared with scripts/assessment_engine.py.

Run: python3 -m unittest scripts/test_fill_workbook.py
"""
from __future__ import annotations

import argparse
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import assessment_engine as eng  # noqa: E402
import fill_workbook as fw  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None
try:
    import formulas
except ImportError:
    formulas = None


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class FillWorkbookTest(unittest.TestCase):

    @classmethod
    def setUpClass(cls) -> None:
        cls.tmp = Path(tempfile.mkdtemp())
        cls.respostas = json.loads(
            (ROOT / "responses.json.example").read_text("utf-8"))
        cls.framework = json.loads(
            (ROOT / "framework.json").read_text("utf-8"))
        args = argparse.Namespace(
            respostas=str(ROOT / "responses.json.example"),
            out=str(cls.tmp))
        assert fw.run(args) == 0
        cls.path = next(cls.tmp.glob("scoring-v1-*.xlsx"))
        cls.wb = openpyxl.load_workbook(cls.path)
        cls.names = fw.HEADERS[fw.locale_of(cls.respostas)]

    def test_answers_sheet_has_every_question(self) -> None:
        ws = self.wb[self.names["answers"]]
        qids = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        expected = [q["id"] for p in self.framework["pillars"]
                    for c in p["capabilities"] for q in c["questions"]]
        self.assertEqual(qids, expected)
        for r in range(2, ws.max_row + 1):
            level = self.respostas["responses"][ws.cell(r, 1).value]["level"]
            self.assertEqual(ws.cell(r, 5).value, level)

    def test_teaching_sheet_uses_framework_weights(self) -> None:
        weights = {q["id"]: q["weight"] for p in self.framework["pillars"]
                   for c in p["capabilities"] for q in c["questions"]}
        ws = self.wb["Exemplo P1 (P1-C1)"]
        for r in range(5, 10):
            self.assertEqual(ws.cell(r, 5).value, weights[ws.cell(r, 1).value])

    def test_text_cells_are_not_formulas(self) -> None:
        ws = self.wb.worksheets[0]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and \
                        cell.value.startswith("= "):
                    self.assertEqual(cell.data_type, "s", cell.coordinate)

    @unittest.skipIf(formulas is None, "optional 'formulas' not installed")
    def test_formulas_match_engine(self) -> None:
        sol = formulas.ExcelModel().loads(str(self.path)).finish()
        sol = sol.calculate()
        name = self.path.name
        scores = eng.compute_scores(self.framework, self.respostas)

        def value(sheet: str, cell: str):
            return sol[f"'[{name}]{sheet.upper()}'!{cell}"].value[0][0]

        for i, cap in enumerate(scores["capabilities"], start=2):
            got = value(self.names["caps"], f"H{i}")
            if cap["score"] is None:
                self.assertEqual(got, "")
            else:
                self.assertAlmostEqual(got, cap["score"], places=3)
        overall = value(self.names["summary"],
                        f"C{len(self.framework['pillars']) + 3}")
        self.assertAlmostEqual(overall, scores["overall"]["score"], places=3)


if __name__ == "__main__":
    unittest.main()
